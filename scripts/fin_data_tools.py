#!/usr/bin/env python3
"""
fin_data_tools.py — Per-file data tools that the agent can call.

Subcommands (all output JSON to stdout):
  excel-list    --path FILE.xlsx                  → list sheet names
  excel-read    --path FILE.xlsx [--sheet NAME]   → dump sheet as 2D array
  pdf-read      --path FILE.pdf                   → extract text
  chart-bar     --x LABEL,LBL,... --y 1.0,2.0,...
                   --output PNG                   → render bar chart
  chart-line    --x LBL,.. --series "name=1,2,3|name2=2,3,4"
                   --output PNG                   → render line chart
  report-build  --uploads DIR [--out DIR] [--charts DIR]
                   [--company NAME] [--quarter 2025Q2] → build full bundle
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from pathlib import Path

def _json_out(obj: dict) -> None:
    sys.stdout.write(json.dumps(obj, indent=2, ensure_ascii=False))
    sys.stdout.write("\n")

def cmd_excel_list(args) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(args.path, read_only=True, data_only=True)
    _json_out({"ok": True, "path": args.path, "sheets": [w.title for w in wb.worksheets]})
    return 0

def cmd_excel_read(args) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(args.path, data_only=True, read_only=True)
    target = args.sheet or wb.worksheets[0].title
    if target not in [w.title for w in wb.worksheets]:
        _json_out({"ok": False, "error": f"sheet '{target}' not found"})
        return 2
    ws = wb[target]
    rows: list[list[str]] = []
    for r in ws.iter_rows(values_only=True):
        row = [("" if c is None else str(c)) for c in r]
        if any(c.strip() for c in row):
            rows.append(row)
    _json_out({"ok": True, "path": args.path, "sheet": target, "rows": rows})
    return 0

def cmd_pdf_read(args) -> int:
    import pdfplumber
    pages = []
    with pdfplumber.open(args.path) as pdf:
        for i, page in enumerate(pdf.pages):
            pages.append({"page": i + 1, "text": page.extract_text() or ""})
    _json_out({"ok": True, "path": args.path, "pages": pages, "full_text": "\n".join(p["text"] for p in pages)})
    return 0

def _ensure_parent(p: str) -> None:
    Path(p).parent.mkdir(parents=True, exist_ok=True)

def cmd_chart_bar(args) -> int:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    xlabels = [s.strip() for s in args.x.split(",") if s.strip()]
    yvals   = [float(s.strip()) for s in args.y.split(",") if s.strip()]
    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(xlabels, yvals, color="#1f77b4", edgecolor="white")
    for i, v in enumerate(yvals):
        ax.text(i, v + max(yvals) * 0.01, f"{v:g}", ha="center", fontsize=9, fontweight="bold")
    ax.set_title(args.title or "Bar chart")
    ax.set_ylabel(args.ylabel or "Value")
    ax.grid(axis="y", alpha=0.3, linestyle="--")
    ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False)
    fig.tight_layout()
    _ensure_parent(args.output)
    fig.savefig(args.output, dpi=140)
    plt.close(fig)
    _json_out({"ok": True, "output": args.output, "type": "bar", "n": len(xlabels)})
    return 0

def cmd_chart_line(args) -> int:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    xlabels = [s.strip() for s in args.x.split(",") if s.strip()]
    series_raw = [s for s in args.series.split("|") if s.strip()]
    series = []
    for s in series_raw:
        if "=" not in s:
            _json_out({"ok": False, "error": f"series '{s}' missing name=values"})
            return 2
        name, vals = s.split("=", 1)
        vs = [float(v.strip()) for v in vals.split(",") if v.strip()]
        series.append((name.strip(), vs))
    fig, ax = plt.subplots(figsize=(8, 4.5))
    for i, (n, v) in enumerate(series):
        ax.plot(xlabels, v, marker="o", linewidth=2, label=n)
    ax.set_title(args.title or "Line chart")
    ax.set_ylabel(args.ylabel or "Value")
    ax.legend()
    ax.grid(axis="y", alpha=0.3, linestyle="--")
    ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False)
    fig.tight_layout()
    _ensure_parent(args.output)
    fig.savefig(args.output, dpi=140)
    plt.close(fig)
    _json_out({"ok": True, "output": args.output, "type": "line", "n_series": len(series)})
    return 0

def cmd_report_build(args) -> int:
    # Delegate to fin_report.py via a subprocess to keep concerns separated
    import subprocess
    script = Path(__file__).parent / "fin_report.py"
    if not script.exists():
        _json_out({"ok": False, "error": f"fin_report.py not found at {script}"})
        return 2
    cmd = [
        "python3", str(script),
        "--uploads", args.uploads,
        "--out",     args.out,
        "--charts",  args.charts,
        "--company", args.company,
    ]
    if args.quarter:
        cmd += ["--quarter", args.quarter]
    proc = subprocess.run(cmd, capture_output=True, text=True)
    if proc.returncode != 0:
        _json_out({"ok": False, "error": proc.stderr or proc.stdout, "returncode": proc.returncode})
        return proc.returncode
    # fin_report.py prints the summary as JSON — surface it directly
    sys.stdout.write(proc.stdout)
    return 0

def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("excel-list")
    p.add_argument("--path", required=True); p.set_defaults(fn=cmd_excel_list)

    p = sub.add_parser("excel-read")
    p.add_argument("--path", required=True)
    p.add_argument("--sheet", default=None); p.set_defaults(fn=cmd_excel_read)

    p = sub.add_parser("pdf-read")
    p.add_argument("--path", required=True); p.set_defaults(fn=cmd_pdf_read)

    p = sub.add_parser("chart-bar")
    p.add_argument("--x",     required=True, help="comma-separated labels")
    p.add_argument("--y",     required=True, help="comma-separated values")
    p.add_argument("--output", required=True)
    p.add_argument("--title",  default=None)
    p.add_argument("--ylabel", default=None); p.set_defaults(fn=cmd_chart_bar)

    p = sub.add_parser("chart-line")
    p.add_argument("--x",     required=True)
    p.add_argument("--series", required=True, help="name=v1,v2,v3|name2=v1,v2")
    p.add_argument("--output", required=True)
    p.add_argument("--title",  default=None)
    p.add_argument("--ylabel", default=None); p.set_defaults(fn=cmd_chart_line)

    p = sub.add_parser("report-build")
    p.add_argument("--uploads", required=True)
    p.add_argument("--out",     default="e2e-output/fin-report/reports")
    p.add_argument("--charts",  default="e2e-output/fin-report/charts")
    p.add_argument("--company", default="Globex Technologies Inc.")
    p.add_argument("--quarter", default=None); p.set_defaults(fn=cmd_report_build)

    args = ap.parse_args()
    return args.fn(args)

if __name__ == "__main__":
    sys.exit(main())
